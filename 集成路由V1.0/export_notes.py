#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 PPTX 的**旁白**与**页面说明**结构化导出为与 PPTX 同级的 Markdown。

两份产物（均与 PPTX 同级）：

    ① <同名>.md          讲稿：逐页「标题 + 旁白」（喂 TTS / 配音 / 提词）
    ② <同名>_页面说明.md  页面说明：逐页「核心信息 / 画面构成 / 层次关系 /
                          强调线索 / 页面文字」（喂下游 AI，告诉它重点在哪、哪里该高亮）

旁白来源优先级：
    ① PPTX 自带的演讲者备注（最权威：就是这份片子实际会念的内容）
    ② skill 项目里的 notes/total.md（备注没嵌进去时的兜底）
页标题来源优先级：
    ① 同级/子级的 design_spec.md 的 §IX（``#### Slide NN - 标题``）
    ② PPTX 里字号最大的那段文字
    ③ ``第 N 页``
页面说明来源（**纯抽取，不再调 LLM**）：
    · design_spec.md §IX 的 Core message / Composition / Relationships / Content
    · spec_lock.md 的配色与字体（告诉下游「强调色该用哪个」）
    · svg_output/*.svg 的实际可见文字（对齐真实成品）

用法::

    python export_notes.py <pptx>                 # 两份都导
    python export_notes.py <运行目录>              # 目录下每个 pptx 各导一份
    python export_notes.py <pptx> --only notes     # 只导讲稿（或 pages）
    python export_notes.py <pptx> -o 讲稿.md       # 单文件时指定讲稿输出路径
    python export_notes.py <pptx> --topic 圆锥曲线解题技巧

产物形如::

    # 圆锥曲线解题技巧 · 讲稿

    > 来源：圆锥曲线解题技巧_195112.pptx
    > 页数：12　|　有旁白：12 页　|　导出：2026-10-18 11:25

    ---

    ## 第 1 页 · 圆锥曲线解题技巧（封面）

    同学们好，今天这节课……
"""
from __future__ import annotations

import argparse
import html
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

#: design_spec.md 里的页清单标题：`#### Slide 01 - 圆锥曲线解题技巧（封面）`
_SLIDE_TITLE_RE = re.compile(r"^####\s*Slide\s*(\d{1,3})\s*[-–—:：]\s*(.+?)\s*$", re.M)
#: PPTX 文件名里的时间戳后缀：`圆锥曲线解题技巧_195112`
_STAMP_SUFFIX_RE = re.compile(r"_\d{6}$")
#: 备注里常见的前缀噪声
_NOTE_PREFIX_RE = re.compile(r"^(?:旁白|讲稿|讲者备注|备注)\s*[：:]\s*")
#: design_spec.md 的页块：`#### Slide 01 - 标题`
_SLIDE_BLOCK_RE = re.compile(r"^####\s*Slide\s*(\d{1,3})\s*[-–—:：]\s*(.+?)\s*$", re.M)
#: 页块里的字段行：`- **Core message**: ...`
_FIELD_RE = re.compile(r"^-\s*\*\*(.+?)\*\*\s*[：:]\s*(.*)$", re.M)
#: 部分标题：`### Part 1: ...`
_PART_RE = re.compile(r"^###\s+(Part\s*\d+[^\n]*)$", re.M)
#: 分句
_SENT_SPLIT_RE = re.compile(r"[。；;]+")
#: 视觉强调线索词
_EMPHASIS_RE = re.compile(r"高亮|强调|突出|加粗|主色|对比色|浅蓝|浅灰|暖色|冷色|色块|斜标|标注|放大|引导线")
#: 字段英文名 → 中文标签
_FIELD_LABELS = {
    "Audience move": "认知目标",
    "Relationships": "层次关系",
    "Composition": "画面构成",
    "Core message": "核心信息",
    "Content": "页面内容",
    "Mathematical content": "数学内容",
}
_VISUAL_FIELDS = ("Composition", "Relationships", "Content", "Core message")


# ---------------------------------------------------------------------------
# 读取：旁白
# ---------------------------------------------------------------------------
def read_notes_from_pptx(pptx_path: Path) -> List[str]:
    """按页序读出 PPTX 里的演讲者备注；缺失的页给空串。"""
    try:
        from pptx import Presentation
    except ModuleNotFoundError as exc:  # pragma: no cover
        raise RuntimeError("缺少 python-pptx：pip install python-pptx") from exc

    prs = Presentation(str(pptx_path))
    out: List[str] = []
    for slide in prs.slides:
        text = ""
        if slide.has_notes_slide:
            try:
                text = slide.notes_slide.notes_text_frame.text or ""
            except (AttributeError, KeyError):
                text = ""
        out.append(_clean_note(text))
    return out


def _clean_note(text: str) -> str:
    """去掉备注里的前后空白与「旁白：」这类前缀，保留正文换行。"""
    lines = [ln.rstrip() for ln in (text or "").replace("\r\n", "\n").split("\n")]
    while lines and not lines[0].strip():
        lines.pop(0)
    while lines and not lines[-1].strip():
        lines.pop()
    body = "\n".join(lines)
    return _NOTE_PREFIX_RE.sub("", body, count=1).strip()


def find_in_project(pptx_path: Path, rel_glob: str, max_depth: int = 3) -> Optional[Path]:
    """在 pptx 同级目录及若干层子目录里找第一个匹配文件（skill 项目文件夹就在旁边）。"""
    base = pptx_path.parent
    direct = base / rel_glob
    if direct.is_file():
        return direct
    hits = sorted(base.glob(f"*/{rel_glob}")) + sorted(base.glob(f"*/*/{rel_glob}"))
    hits = [p for p in hits if p.is_file()]
    if hits:
        return hits[0]
    for child in sorted(base.rglob(Path(rel_glob).name)):
        if child.is_file() and len(child.relative_to(base).parts) <= max_depth + 1:
            return child
    return None


def read_notes_from_total_md(pptx_path: Path) -> Optional[List[str]]:
    """兜底：从 skill 项目的 notes/total.md 取旁白（`# stem` + `---` 分隔）。"""
    f = find_in_project(pptx_path, "notes/total.md")
    if not f:
        return None
    chunks = re.split(r"\n\s*---\s*\n", f.read_text(encoding="utf-8", errors="replace"))
    notes: List[str] = []
    for chunk in chunks:
        body = re.sub(r"^\s*#\s*[^\n]*\n", "", chunk, count=1)
        notes.append(body.strip())
    return notes if any(notes) else None


# ---------------------------------------------------------------------------
# 读取：页标题
# ---------------------------------------------------------------------------
def read_titles_from_design_spec(pptx_path: Path) -> Dict[int, str]:
    """从 design_spec.md 的 §IX 取 ``Slide NN → 标题`` 映射。"""
    f = find_in_project(pptx_path, "design_spec.md")
    if not f:
        return {}
    text = f.read_text(encoding="utf-8", errors="replace")
    return {int(num): title.strip() for num, title in _SLIDE_TITLE_RE.findall(text)}


def _iter_text_frames(shapes):
    """递归遍历（含组合形状）里的文本框。"""
    for sh in shapes:
        if getattr(sh, "shape_type", None) == 6 or hasattr(sh, "shapes"):
            try:
                yield from _iter_text_frames(sh.shapes)
                continue
            except (AttributeError, TypeError):
                pass
        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            yield sh


def read_titles_from_pptx(pptx_path: Path) -> Dict[int, str]:
    """兜底：没有 design_spec 时，逐页取字号最大的那段文字当标题。"""
    try:
        from pptx import Presentation
    except ModuleNotFoundError:  # pragma: no cover
        return {}
    prs = Presentation(str(pptx_path))
    return {i: guess_title_from_slide(slide)
            for i, slide in enumerate(prs.slides, 1)}


def guess_title_from_slide(slide) -> str:
    """兜底：取该页字号最大的那段文字（同字号取更长的那段）。"""
    best_text, best_score = "", -1.0
    for sh in _iter_text_frames(slide.shapes):
        frame = sh.text_frame
        for para in frame.paragraphs:
            text = "".join(r.text for r in para.runs).strip()
            if not text:
                continue
            sizes = [r.font.size.pt for r in para.runs
                     if r.font.size is not None]
            size = max(sizes) if sizes else 0.0
            score = size * 100 + len(text)
            if score > best_score:
                best_text, best_score = text, score
    return best_text


# ---------------------------------------------------------------------------
# 组装与导出
# ---------------------------------------------------------------------------
def build_markdown(pptx_path: Path, topic: Optional[str] = None,
                   notes: Optional[List[str]] = None,
                   titles: Optional[Dict[int, str]] = None) -> str:
    """把「逐页标题 + 旁白」拼成结构化 Markdown。"""
    if notes is None:
        notes = read_notes_from_pptx(pptx_path)
    if titles is None:
        titles = read_titles_from_design_spec(pptx_path)
    missing = [i for i in range(1, len(notes) + 1) if not (titles.get(i) or "").strip()]
    if missing:  # design_spec 缺失/不全 → 用字号最大的文字补齐
        fallback = read_titles_from_pptx(pptx_path)
        titles = {i: (titles.get(i) or fallback.get(i) or "") for i in range(1, len(notes) + 1)}

    topic = topic or _STAMP_SUFFIX_RE.sub("", pptx_path.stem)
    filled = sum(1 for n in notes if n.strip())

    lines = [
        f"# {topic} · 讲稿",
        "",
        f"> 来源：{pptx_path.name}  ",
        f"> 页数：{len(notes)}　|　有旁白：{filled} 页　|　导出：{datetime.now():%Y-%m-%d %H:%M}",
        "",
        "---",
        "",
    ]
    for i, note in enumerate(notes, 1):
        title = (titles.get(i) or "").strip()
        lines += [f"## 第 {i} 页 · {title}" if title else f"## 第 {i} 页", ""]
        lines += [note.strip() if note.strip() else "_（本页无旁白）_", ""]
        lines += ["---", ""]

    while lines and lines[-1] == "":
        lines.pop()
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# 页面说明：内容重点 + 视觉构成
# ---------------------------------------------------------------------------
def find_dir_in_project(pptx_path: Path, name: str, max_depth: int = 3) -> Optional[Path]:
    """在 pptx 同级及若干层子目录里找目录（skill 项目就在旁边）。"""
    base = pptx_path.parent
    if (base / name).is_dir():
        return base / name
    for p in sorted(base.iterdir()):
        if p.is_dir() and (p / name).is_dir():
            return p / name
    for p in sorted(base.rglob(name)):
        if p.is_dir() and len(p.relative_to(base).parts) <= max_depth:
            return p
    return None


def read_slides_from_design_spec(pptx_path: Path) -> List[dict]:
    """从 design_spec.md 的 §IX 解析逐页字段（标题/核心信息/画面构成/层次关系/内容）。"""
    f = find_in_project(pptx_path, "design_spec.md")
    if not f:
        return []
    text = f.read_text(encoding="utf-8", errors="replace")
    marks = list(_SLIDE_BLOCK_RE.finditer(text))
    parts = [(m.start(), m.group(1).strip()) for m in _PART_RE.finditer(text)]
    slides: List[dict] = []
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        body = text[m.end():end]
        part = ""
        for pos, name in parts:
            if pos < m.start():
                part = name
            else:
                break
        fields = {fm.group(1).strip(): fm.group(2).strip()
                  for fm in _FIELD_RE.finditer(body)}
        slides.append({"index": int(m.group(1)), "title": m.group(2).strip(),
                       "part": part, "fields": fields})
    return slides


def read_page_texts(pptx_path: Path) -> Dict[int, List[str]]:
    """从 svg_output/*.svg 取每页**实际**可见文字（按文件名数字前缀对号）。"""
    d = find_dir_in_project(pptx_path, "svg_output")
    if not d:
        return {}
    out: Dict[int, List[str]] = {}
    for svg in sorted(d.glob("*.svg")):
        m = re.match(r"(\d+)", svg.stem)
        if not m:
            continue
        raw = svg.read_text(encoding="utf-8", errors="replace")
        texts: List[str] = []
        for chunk in re.findall(r"<text[^>]*>(.*?)</text>", raw, re.S):
            s = html.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", chunk))).strip()
            if not s:
                continue
            if len(s) == 1 and not re.search(r"[\u4e00-\u9fff]", s):
                continue  # 单字标签（A / B / ?）不算“重点”
            if s not in texts:
                texts.append(s)
        out[int(m.group(1))] = texts
    return out


def read_visual_baseline(pptx_path: Path) -> Dict[str, str]:
    """从 spec_lock.md 取全量视觉基线（画布/配色/字体），供「高亮该用什么色」参考。"""
    f = find_in_project(pptx_path, "spec_lock.md")
    if not f:
        return {}
    out: Dict[str, str] = {}
    for line in f.read_text(encoding="utf-8", errors="replace").splitlines():
        m = re.match(r"^-\s*([A-Za-z_]+)\s*:\s*(.+)$", line.strip())
        if m:
            out.setdefault(m.group(1), m.group(2).strip())
    return out


def emphasis_cues(fields: Dict[str, str]) -> List[str]:
    """从视觉字段里摘出讲「哪里该高亮/突出」的原句（只引用，不生成）。"""
    cues: List[str] = []
    for key in _VISUAL_FIELDS:
        for sent in _SENT_SPLIT_RE.split(fields.get(key, "") or ""):
            sent = sent.strip(" ，,、")
            if sent and _EMPHASIS_RE.search(sent) and sent not in cues:
                cues.append(sent)
    return cues


def build_pages_markdown(pptx_path: Path, topic: Optional[str] = None,
                         slides: Optional[List[dict]] = None) -> str:
    """拼「页面说明」：每页讲什么 + 画面怎么排 + 哪里该高亮。"""
    if slides is None:
        slides = read_slides_from_design_spec(pptx_path)
    texts = read_page_texts(pptx_path)
    base = read_visual_baseline(pptx_path)
    topic = topic or _STAMP_SUFFIX_RE.sub("", pptx_path.stem)

    lines = [f"# {topic} · 页面说明", ""]
    lines.append(f"> 来源：{pptx_path.name}  ")
    lines.append(f"> 页数：{len(slides)}　|　spec：design_spec.md　|　导出：{datetime.now():%Y-%m-%d %H:%M}")
    lines.append(">")
    lines.append("> 用途：给下游 AI 读——每页**讲什么**（重点）、**怎么排**（视觉）、**哪里该高亮**。")
    lines.append("")
    if base:
        style = "　".join(x for x in (
            base.get("format") or (f"viewBox {base['viewBox']}" if base.get("viewBox") else ""),
            base.get("visual_style") or "",
            (f"主色 {base['primary']}" if base.get("primary") else ""),
            (f"强调色 {base['accent']}" if base.get("accent") else ""),
            (f"字体 {base['font_family']}" if base.get("font_family") else ""),
        ) if x)
        if style:
            lines += [f"**视觉基线**：{style}", ""]
    lines += ["---", ""]

    for slide in slides:
        idx, title, fields = slide["index"], slide["title"], slide["fields"]
        lines.append(f"## 第 {idx} 页 · {title}")
        lines.append("")
        if slide.get("part"):
            lines.append(f"- **所属部分**：{slide['part']}")
        for key in ("Core message", "Composition", "Relationships", "Content",
                    "Mathematical content", "Audience move"):
            value = (fields.get(key) or "").strip()
            if value and value not in ("—", "-"):
                lines.append(f"- **{_FIELD_LABELS[key]}**：{value}")
        cues = emphasis_cues(fields)
        if cues:
            lines.append("- **强调线索**：")
            lines += [f"  - {c}" for c in cues]
        if texts.get(idx):
            lines.append(f"- **页面文字**：{' / '.join(texts[idx])}")
        lines += ["", "---", ""]

    while lines and lines[-1] == "":
        lines.pop()
    return "\n".join(lines) + "\n"


def export_pages_md(pptx_path: Path, out_path: Optional[Path] = None,
                    topic: Optional[str] = None, quiet: bool = False) -> Optional[Path]:
    """导出「页面说明」Markdown（内容重点 + 视觉构成 + 高亮线索）。"""
    pptx_path = Path(pptx_path)
    if not pptx_path.is_file():
        raise FileNotFoundError(f"找不到 PPTX：{pptx_path}")
    slides = read_slides_from_design_spec(pptx_path)
    if not slides:
        if not quiet:
            print(f"   ⚠️  {pptx_path.name}：没找到 design_spec.md，跳过页面说明导出")
        return None
    dst = Path(out_path) if out_path else pptx_path.with_name(f"{pptx_path.stem}_页面说明.md")
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text(build_pages_markdown(pptx_path, topic=topic, slides=slides),
                   encoding="utf-8")
    if not quiet:
        cues = sum(1 for s in slides if emphasis_cues(s["fields"]))
        print(f"   ✓ 页面说明已导出：{dst}（{len(slides)} 页，{cues} 页含强调线索）")
    return dst


def export_narration_md(pptx_path: Path, out_path: Optional[Path] = None,
                        topic: Optional[str] = None,
                        quiet: bool = False) -> Optional[Path]:
    """导出旁白 Markdown；PPTX 无备注且找不到 total.md 时返回 None。"""
    pptx_path = Path(pptx_path)
    if not pptx_path.is_file():
        raise FileNotFoundError(f"找不到 PPTX：{pptx_path}")

    notes = read_notes_from_pptx(pptx_path)
    source = "PPTX 备注"
    if not any(n.strip() for n in notes):
        fallback = read_notes_from_total_md(pptx_path)
        if fallback:
            notes, source = fallback, "notes/total.md"
        else:
            if not quiet:
                print(f"   ⚠️  {pptx_path.name}：备注为空，也没找到 notes/total.md，跳过讲稿导出")
            return None

    dst = Path(out_path) if out_path else pptx_path.with_suffix(".md")
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text(build_markdown(pptx_path, topic=topic, notes=notes),
                   encoding="utf-8")
    if not quiet:
        filled = sum(1 for n in notes if n.strip())
        print(f"   ✓ 讲稿已导出：{dst}（{filled}/{len(notes)} 页，来源={source}）")
    return dst


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def _iter_pptx(target: Path) -> List[Path]:
    """目录只取**顶层** pptx：运行目录里那是成品，skill 项目里的 exports/ 是中间产物。"""
    if target.is_file() and target.suffix.lower() == ".pptx":
        return [target]
    if target.is_dir():
        top = sorted(p for p in target.glob("*.pptx") if not p.name.startswith("~$"))
        if top:
            return top
        return sorted(p for p in target.rglob("*.pptx")
                      if not p.name.startswith("~$"))
    return []


def main() -> int:
    ap = argparse.ArgumentParser(
        description="把 PPTX 的旁白（演讲者备注）结构化导出为同级 Markdown")
    ap.add_argument("target", help="PPTX 文件，或包含 PPTX 的运行目录")
    ap.add_argument("-o", "--output", default=None, help="讲稿输出路径（仅单文件时有效）")
    ap.add_argument("--topic", default=None, help="标题里的主题名（默认取文件名去掉时间戳）")
    ap.add_argument("--only", choices=("all", "notes", "pages"), default="all",
                    help="只导某一类产物（默认两份都导）")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args()

    target = Path(args.target).expanduser()
    if not target.exists():
        print(f"✗ 路径不存在：{target}")
        return 2

    files = _iter_pptx(target)
    if not files:
        print(f"✗ 没找到 .pptx：{target}")
        return 2

    done = 0
    for pptx in files:
        out = Path(args.output) if (args.output and len(files) == 1) else None
        try:
            if args.only in ("all", "notes"):
                if export_narration_md(pptx, out_path=out, topic=args.topic,
                                       quiet=args.quiet) is not None:
                    done += 1
            if args.only in ("all", "pages"):
                if export_pages_md(pptx, topic=args.topic, quiet=args.quiet) is not None:
                    done += 1
        except (RuntimeError, FileNotFoundError, KeyError) as exc:
            print(f"   ✗ {pptx.name}：{exc}")
    if not args.quiet:
        print(f"\n共处理 {len(files)} 个 PPTX，成功导出 {done} 份产物。")
    return 0 if done else 1


if __name__ == "__main__":
    raise SystemExit(main())
